import filetype
from pypdf import PdfReader
from docx import Document
import re

from nltk.tokenize import word_tokenize
from nltk.corpus import stopwords
from nltk.stem import WordNetLemmatizer
import spacy
import nltk




nlp = spacy.load("en_core_web_sm")

# NLTK data location
nltk.data.path = [
    r"C:\Users\B23F0542AI156\nltk_data"
]

skills_section = [
        "skills",
        "technical skills",
        "professional skills"
    ]

target_skills = [
            "Python",
            "Django",
            "SQL",
            "AWS",
            "React",
            "Javascript",
            "MongoDB",
            "Git",
            "Artificial Intelligence",
            "Data Science",
]
target_skills=[skill.lower() for skill in target_skills]
keywords=["Python","Django","Web-Development","Javascript","Data Science","Artifical Intelligence"]

def count_keywords(keywords,extracted_text):
    total_keywords=0
    for keyword in keywords:
        if keyword in extracted_text:
            extracted_text_list=extracted_text.split()
            count_frequency=extracted_text_list.count(keyword)
            total_keywords+=count_frequency
        average_keyword_score=len(keywords)/len(total_keywords)
        return total_keywords,average_keyword_score
    

def skills_check(target_skills,skills):
    detected_skills = []
    resume_score=""
    for skill in target_skills:
        for resume_text in skills:
         if skill in resume_text:
            detected_skills.append(skill)
    # Counting total detected skills
    detected_skills_percentage=len(detected_skills)/len(target_skills)*100
    if detected_skills_percentage<80:
        resume_score=
        return "The skills matching score with job description is below 80%"
    
    return detected_skills,detected_skills_percentage
    




    


    

    

    

        

        
        

                

    return detected_skills



def list_missing_skills(target_skills,detected_skills):
    target_skills = [
            "Python",
            "Django",
            "SQL",
            "AWS",
            "React",
            "Javascript",
            "MongoDB",
            "Git",
            "Artificial Intelligence",
            "Data Science",
            "Azure"
        ]
    if target_skills not in detected_skills:
        return list(target_skills)




file_path = input("Enter your File Path: ")

try:

    check_file_path = filetype.guess(file_path)

    if check_file_path is None:
        print("Not a valid file")

    else:

        print(check_file_path)

        extension = check_file_path.extension
        mime = check_file_path.mime

       

        if extension == "pdf":

            reader = PdfReader(file_path)

            extracted_text = ""

            for page in reader.pages:
                page_text = page.extract_text()

                if page_text:
                    extracted_text += page_text + "\n"

            print(extracted_text)

        elif extension == "docx":

            get_content = Document(file_path)

            detected_skills = []
            

            skills_section = [
        "skills",
        "technical skills",
        "professional skills"
    ]
            inside_skills=False
            skills=[]
            text=""
        for paragraph in get_content.paragraphs:
             
            if paragraph.style.name.startswith("Heading 2") and paragraph.text.lower() in skills_section:
             inside_skills=True
             continue
            if inside_skills:
                if paragraph.style.name.startswith("Heading 2"):
                    break
                skills.append(paragraph.text)
        skills=[skill.lower() for skill in skills]
        counting_matched_skills=skills_check(target_skills,skills)
        print(counting_matched_skills)

        print(skills)

            

        # text+=paragraph.text
        # text=text.lower()
        # find_skills=re.search(r"(skills|technical skills|professional skills):*\s+(.*?)(?=\n[a-zA-Z][a-zA-Z]*\s+)",text).group()
        # print(find_skills)

            


    #         if paragraph.style.name.startswith("Heading 1") and heading in skills_section:
    #             inside_skills = True
    #             continue

    #         if inside_skills:

    #             if paragraph.style.name.startswith("Heading 1"):
    #                 inside_skills = False
    #                 continue

    #         text += paragraph.text + " "

    # text = text.lower()

    # print(text)
            # calling skills check function
            # skills_function=skills_check(detected_skills,text)
            # print(skills_function)
            
            #         text+= paragraph.text.strip()
            # text=re.sub(r"[^a-zA-Z\d+@]"," ",text)
            # cleaned_text=text
            # cleaned_text=cleaned_text.lower()
            # show_keywords=count_keywords(keywords,cleaned_text)
            # print(show_keywords)
            
#             # calling a skills check function
#             detected_skills_keywords="skills_check"(get_content)
#             # Comparing detected skills against actual skills list
#             list_missing_skills(target_skills,skills_check())
#             # Designing a scoring function
#             #list_missing_skills(target_skills,skills_check())
#             # Counting a specific keyword
            

#     print("Detected Skills:", detected_skills_keywords)

#     # Comparing Detected Skills
    



        
           #Designing Rubric System
# #             skills_percentage=40
# #             experience=30
# #             keywords=30
# #             calculate_score=len(detected_skills)/len(skills)*100
# #             print(calculate_score)

# #             # Getting Experience
# #             experience_section=["Professional Experience",
# #                                 "Technical Experience","Experience"]
# #             experience_section=[experience.lower() for experience in experience_section]
# #             for i,paragraph in enumerate(paragraphs):
# #                 if (paragraph.style.name.startswith("Heading 1") and paragraph.text in experience_section):
# #                     for get_experience in paragraph[i+1:]:
# #                         if get_experience.style.name.startswith("Heading 1"):
# #                             break
# #                         record_experience=get_experience.text
# #                         experience_years=nlp(record_experience)
# #                         print(experience_years.ents)
                        
                            


# #             text = text.lower().strip()

# #             #print("Resume Text:")
# #             #print(text)

# #             # Counting Words
# #             print(
# #                 "Word Count:",
# #                 len(text.split())
# #             )

           
# #             cleaned_text = re.sub(
# #                 r'[^a-zA-Z0-9\s]',
# #                 '',
# #                 text
# #             )

# #             print("Cleaned Text:")
# #             #print(cleaned_text)

           

# #             tokens = word_tokenize(cleaned_text)

# #             print("Tokens:")
# #             print(tokens)

            

# #             stop_words = set(
# #                 stopwords.words("english")
# #             )

# #             filtered_tokens = [
# #                 word
# #                 for word in tokens
# #                 if word not in stop_words
# #             ]

# #             print("After Stopword Removal:")
# #             print(filtered_tokens)

            

# #             lemmatizer = WordNetLemmatizer()

# #             lemmatized = [
# #                 lemmatizer.lemmatize(word)
# #                 for word in filtered_tokens
# #             ]

# #             print("Lemmatized:")
# #             print(lemmatized)

            

            

# #             doc = nlp(text)

# #             print("Named Entities:")

# #             for entity in doc.ents:

# #                 print(
# #                     entity.text,
# #                     ">",
# #                     entity.label_
# #                 )

        

# #         else:

# #             print(
# #                 "Not supported file type for Resume"
# #             )


except FileNotFoundError:
    pass

#         print(
#        "The path entered does not exist."
#    )


# # 
# def skills_check(get_content):

#     skills = [
#         "Python",
#         "Django",
#         "SQL",
#         "AWS",
#         "React",
#         "Javascript",
#         "MongoDB",
#         "Git",
#         "Artificial Intelligence",
#         "Data Science",
#         "Azure"
#     ]

#     detected_skills = []

#     skills_section = [
#         "skills",
#         "technical skills",
#         "professional skills"
#     ]

#     inside_skills = False

#     for paragraph in get_content.paragraphs:

#         text = paragraph.text.strip()

#         if not text:
#             continue

#         # Check whether this paragraph is a Skills heading
#         if text.lower().rstrip(":") in skills_section:
#             inside_skills = True
#             continue

#         # If we are inside Skills section
#         if inside_skills:

#             # Check if another major heading has started
#             if paragraph.style.name.startswith("Heading 1"):
#                 inside_skills = False
#                 continue

#             # Check skills in this paragraph
#             for skill in skills:

#                 if re.search(
#                     rf"\b{re.escape(skill)}\b",
#                     text,
#                     re.IGNORECASE
#                 ):
#                     if skill not in detected_skills:
#                         detected_skills.append(skill)

#     return detected_skills

# def score_resume(target_skills,detected_skills):
#     detected_skills=skills_check()
#     skills_count=target_skills/detected_skills
#     skills_count=skills_count*100
#     return skills_count

